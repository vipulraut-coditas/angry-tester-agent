#!/usr/bin/env python3
"""Angry Tester report builder.

Reads findings.json + test_results.json (see ../schemas/findings.schema.json) and writes:
  <out>/ANGRY_TESTER_REPORT[_<feature>].md
  <out>/angry_tester_report[_<feature>].xlsx

Usage:
  python build_report.py --findings angry-tester/findings.json \
                         --tests angry-tester/test_results.json \
                         --out angry-tester-report
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

SCHEMA_PATH = Path(__file__).resolve().parent.parent / "schemas" / "findings.schema.json"

SEVERITIES = ["Critical", "High", "Medium", "Low"]
IMPACTS = ["HACK", "DOWN", "CRACK", "BUG"]
CONFIDENCES = ["Confirmed", "Likely", "Theoretical"]
RESULTS = ["PASS", "FAIL", "NOT RUN"]

SEV_FILL = {
    "Critical": ("8B0000", "FFFFFF"),
    "High": ("FF0000", "FFFFFF"),
    "Medium": ("FFA500", "000000"),
    "Low": ("FFFF00", "000000"),
}
RESULT_FILL = {
    "PASS": ("00B050", "FFFFFF"),
    "FAIL": ("FF0000", "FFFFFF"),
    "NOT RUN": ("BFBFBF", "000000"),
}
HEADER_FILL = PatternFill("solid", fgColor="1F1F1F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MAX_COL_WIDTH = 60


# --------------------------------------------------------------------------- validation
def validate(data: dict, def_name: str, label: str) -> list[str]:
    """Validate against the schema. Uses jsonschema if installed, else a minimal check."""
    schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
    try:
        import jsonschema  # type: ignore

        sub = {"$schema": schema["$schema"], "$defs": schema["$defs"], "$ref": f"#/$defs/{def_name}"}
        validator_cls = getattr(jsonschema, "Draft202012Validator", None)
        if validator_cls is None:  # jsonschema < 4.0 has no 2020-12 support
            raise ImportError
        validator = validator_cls(sub)
        return [
            f"{label}: {'/'.join(str(p) for p in e.absolute_path) or '<root>'}: {e.message}"
            for e in sorted(validator.iter_errors(data), key=lambda e: list(e.absolute_path))
        ]
    except ImportError:
        required = schema["$defs"][def_name]["required"]
        return [f"{label}: missing top-level key '{k}'" for k in required if k not in data]


# --------------------------------------------------------------------------- helpers
def sort_findings(findings: list[dict]) -> list[dict]:
    def key(f):
        return (
            SEVERITIES.index(f["severity"]) if f["severity"] in SEVERITIES else 99,
            IMPACTS.index(f["impact"]) if f["impact"] in IMPACTS else 99,
            CONFIDENCES.index(f.get("confidence")) if f.get("confidence") in CONFIDENCES else 99,
            f["id"],
        )

    return sorted(findings, key=key)


def suffix_for(meta: dict) -> str:
    scope = meta.get("scope", {})
    if scope.get("type") == "feature" and scope.get("feature"):
        slug = re.sub(r"[^A-Za-z0-9]+", "_", scope["feature"]).strip("_").lower()
        return f"_{slug}" if slug else ""
    if scope.get("type") == "diff":
        rng = scope.get("diff_range") or "diff"
        slug = re.sub(r"[^A-Za-z0-9]+", "_", rng).strip("_").lower()
        return f"_diff_{slug}"[:60]
    return ""


def scope_line(meta: dict) -> str:
    s = meta["scope"]
    if s["type"] == "full":
        return "Complete application"
    if s["type"] == "feature":
        locs = ", ".join(s.get("locations", [])) or "n/a"
        desc = f" ({s['description']})" if s.get("description") else ""
        return f"Feature: {s.get('feature', '?')}{desc} - locations: {locs}"
    return f"Recent changes: {s.get('diff_range', '?')}"


def md_cell(v) -> str:
    return str(v if v is not None else "").replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def md_table(headers: list[str], rows: list[list]) -> str:
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    out += ["| " + " | ".join(md_cell(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


def fenced(text: str) -> str:
    text = str(text or "").rstrip()
    if "```" in text:  # already contains code fences, keep as is
        return text
    return f"```\n{text}\n```" if text else "_n/a_"


def test_counts(tests: list[dict]) -> Counter:
    return Counter(t["result"] for t in tests)


# --------------------------------------------------------------------------- markdown
def build_markdown(fd: dict, tr: dict) -> str:
    meta, findings, tests = fd["meta"], sort_findings(fd["findings"]), tr["tests"]
    sev = Counter(f["severity"] for f in findings)
    tc = test_counts(tests)
    L: list[str] = []
    title = f"Angry Tester Report - {meta['project']}"
    if meta["scope"]["type"] == "feature":
        title += f" ({meta['scope'].get('feature')})"
    L += [f"# {title}", "", f"_Date: {meta['date']} | Repo: `{meta['repo']}` | Stack: {meta['stack']}_", ""]

    # 1 Scope
    L += ["## 1. Scope", "", f"- **Scope:** {scope_line(meta)}", f"- **Focus:** {', '.join(meta['focus'])}"]
    if meta.get("tools_run"):
        L.append(f"- **Tools run:** {', '.join(meta['tools_run'])}")
    L.append("")

    # 2 Verdict
    L += ["## 2. Verdict", "", f"> {meta['verdict']}", ""]

    # 3 Scorecard
    L += ["## 3. Scorecard", "", "### Findings: severity x impact", ""]
    rows = []
    for s in SEVERITIES:
        row = [f"**{s}**"] + [sum(1 for f in findings if f["severity"] == s and f["impact"] == i) for i in IMPACTS]
        rows.append(row + [sev.get(s, 0)])
    rows.append(["**Total**"] + [sum(1 for f in findings if f["impact"] == i) for i in IMPACTS] + [len(findings)])
    L += [md_table(["Severity"] + IMPACTS + ["Total"], rows), ""]
    L += ["### Tests", ""]
    L += [md_table(["Total", "Run", "Passed", "Failed", "Not run"],
                   [[len(tests), tc["PASS"] + tc["FAIL"], tc["PASS"], tc["FAIL"], tc["NOT RUN"]]]), ""]

    # 4 Top 5
    L += ["## 4. Top 5 ways it gets hacked, goes down or gets cracked", ""]
    verbs = {"HACK": "gets HACKED", "DOWN": "goes DOWN", "CRACK": "gets CRACKED", "BUG": "BREAKS"}
    if findings:
        for n, f in enumerate(findings[:5], 1):
            L.append(f"{n}. **[{f['severity']}] It {verbs.get(f['impact'], f['impact'])}: {f['title']}** "
                     f"(`{f['file']}:{f['line']}`, {f['confidence']}) - {f['scenario']}")
    else:
        L.append("_No findings. Suspicious. Check the Gaps section._")
    L.append("")

    # 5 All findings
    L += ["## 5. All findings", ""]
    L += [md_table(["ID", "Title", "Impact", "Severity", "Confidence", "Scope", "Category", "Location", "Status"],
                   [[f["id"], f["title"], f["impact"], f["severity"], f["confidence"], f["scope"],
                     f["category"], f"`{f['file']}:{f['line']}`", f["status"]] for f in findings])
          if findings else "_None._", ""]

    # 6 Details
    L += ["## 6. Finding details", ""]
    for f in findings:
        L += [f"### {f['id']} - {f['title']}", "",
              f"- **Impact / Severity / Confidence:** {f['impact']} / {f['severity']} / {f['confidence']}",
              f"- **Why this severity:** {f['severity_reason']}",
              f"- **Scope:** {f['scope']}",
              f"- **Category:** {f['category']}" + (f" ({', '.join(f['references'])})" if f.get("references") else ""),
              f"- **Location:** `{f['file']}:{f['line']}`"]
        if f.get("linked_tests"):
            L.append(f"- **Tests:** {', '.join(f['linked_tests'])}")
        L += ["", "**Scenario:** " + f["scenario"], "", "**PoC input:**", "", fenced(f["poc_input"]), "",
              "**Evidence:**", "", fenced(f["evidence"]), "", "**Business impact:** " + f["business_impact"], "",
              "**Fix:**", "", f["fix"], ""]

    # 7 Tests
    L += ["## 7. Test scenarios", ""]
    if tr.get("command"):
        L += [f"Framework: **{tr['framework']}** - command: `{tr['command']}`", ""]
    if tr.get("test_files"):
        L += ["Test files:", ""] + [f"- `{p}`" for p in tr["test_files"]] + [""]
    L += [md_table(["Test ID", "Category", "Target", "Scenario", "Input", "Expected", "Actual", "Result", "Finding"],
                   [[t["test_id"], t["category"], t["target"], t["scenario"], t["input"], t["expected"],
                     t["actual"] + (f" (NOT RUN: {t['not_run_reason']})" if t.get("not_run_reason") else ""),
                     ("✅ " if t["result"] == "PASS" else "❌ " if t["result"] == "FAIL" else "⏸ ") + t["result"],
                     t.get("linked_finding", "")] for t in tests]) if tests else "_No tests._", ""]

    # 8 Trade-offs
    L += ["## 8. Trade-offs", ""]
    L += [md_table(["Area", "Option A", "Option B", "Risk", "Recommendation"],
                   [[t["area"], t["option_a"], t["option_b"], t["risk"], t["recommendation"]] for t in fd["trade_offs"]])
          if fd["trade_offs"] else "_None._", ""]

    # 9 Attack surface
    L += ["## 9. Attack surface", ""]
    L += [md_table(["Entry point", "Method", "Auth", "Roles", "Inputs", "Touches", "Risk notes"],
                   [[a["entry_point"], a["method"], a["auth"], a["roles"], a["inputs"], a["touches"], a["risk_notes"]]
                    for a in fd["attack_surface"]]) if fd["attack_surface"] else "_None recorded._", ""]

    # 10 Gaps
    L += ["## 10. Gaps and what was not tested", ""]
    not_run = [t for t in tests if t["result"] == "NOT RUN"]
    L += [f"- {g}" for g in fd["gaps"]] or ["- _None declared._"]
    L += [f"- Test {t['test_id']} NOT RUN: {t.get('not_run_reason', 'no reason given')}" for t in not_run]
    L.append("")
    return "\n".join(L)


# --------------------------------------------------------------------------- excel
def style_sheet(ws, wrap_cols_max: int = MAX_COL_WIDTH) -> None:
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
    for col in ws.iter_cols(min_row=1, max_row=ws.max_row):
        longest = 0
        for c in col:
            if c.value is not None:
                longest = max(longest, max(len(line) for line in str(c.value).split("\n")))
            if c.row > 1:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 2, 10), wrap_cols_max)


def colour(cell, mapping: dict) -> None:
    if cell.value in mapping:
        bg, fg = mapping[cell.value]
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(bold=True, color=fg)


def xl(v):
    """openpyxl rejects control characters; also cap at Excel's cell limit."""
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return v
    s = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(v))
    if s.startswith(("=", "+", "-", "@")):  # avoid formula injection from PoC payloads
        s = "'" + s
    return s[:32000]


def build_excel(fd: dict, tr: dict, path: Path) -> None:
    meta, findings, tests = fd["meta"], sort_findings(fd["findings"]), tr["tests"]
    sev, imp, tc = Counter(f["severity"] for f in findings), Counter(f["impact"] for f in findings), test_counts(tests)
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Field", "Value"])
    rows = [
        ("Project", meta["project"]), ("Date", meta["date"]), ("Repo", meta["repo"]),
        ("Scope", scope_line(meta)), ("Focus", ", ".join(meta["focus"])), ("Stack", meta["stack"]),
        ("Verdict", meta["verdict"]), ("Total findings", len(findings)),
    ]
    rows += [(f"Severity: {s}", sev.get(s, 0)) for s in SEVERITIES]
    rows += [(f"Impact: {i}", imp.get(i, 0)) for i in IMPACTS]
    rows += [("Tests total", len(tests)), ("Tests run", tc["PASS"] + tc["FAIL"]), ("Tests passed", tc["PASS"]),
             ("Tests failed", tc["FAIL"]), ("Tests not run", tc["NOT RUN"])]
    for r in rows:
        ws.append([xl(c) for c in r])
    style_sheet(ws, 100)
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value).startswith("Severity: "):
            sname = row[0].value.split(": ", 1)[1]
            bg, fg = SEV_FILL[sname]
            row[0].fill = PatternFill("solid", fgColor=bg)
            row[0].font = Font(bold=True, color=fg)

    ws = wb.create_sheet("Findings")
    hdr = ["ID", "Title", "Impact", "Severity", "Confidence", "Scope", "Category", "File", "Line", "Scenario",
           "PoC Input", "Evidence", "Business Impact", "Fix", "Status"]
    ws.append(hdr)
    for f in findings:
        ws.append([xl(v) for v in (f["id"], f["title"], f["impact"], f["severity"], f["confidence"], f["scope"],
                                   f["category"], f["file"], f["line"], f["scenario"], f["poc_input"], f["evidence"],
                                   f["business_impact"], f["fix"], f["status"])])
    style_sheet(ws)
    for row in ws.iter_rows(min_row=2):
        colour(row[hdr.index("Severity")], SEV_FILL)

    ws = wb.create_sheet("Test Cases")
    hdr = ["Test ID", "Category", "Endpoint/Function", "Scenario", "Input", "Expected", "Actual", "Result",
           "Linked Finding"]
    ws.append(hdr)
    for t in tests:
        actual = t["actual"] + (f" (NOT RUN: {t['not_run_reason']})" if t.get("not_run_reason") else "")
        ws.append([xl(v) for v in (t["test_id"], t["category"], t["target"], t["scenario"], t["input"],
                                   t["expected"], actual, t["result"], t.get("linked_finding", ""))])
    style_sheet(ws)
    for row in ws.iter_rows(min_row=2):
        colour(row[hdr.index("Result")], RESULT_FILL)

    ws = wb.create_sheet("Attack Surface")
    ws.append(["Entry Point", "Method", "Auth", "Roles", "Inputs", "Touches", "Risk Notes"])
    for a in fd["attack_surface"]:
        ws.append([xl(a[k]) for k in ("entry_point", "method", "auth", "roles", "inputs", "touches", "risk_notes")])
    style_sheet(ws)

    ws = wb.create_sheet("Trade-offs")
    ws.append(["Area", "Option A", "Option B", "Risk", "Recommendation"])
    for t in fd["trade_offs"]:
        ws.append([xl(t[k]) for k in ("area", "option_a", "option_b", "risk", "recommendation")])
    style_sheet(ws)

    wb.save(path)


# --------------------------------------------------------------------------- main
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--findings", default="angry-tester/findings.json")
    ap.add_argument("--tests", default="angry-tester/test_results.json")
    ap.add_argument("--out", default="angry-tester-report")
    ap.add_argument("--no-validate", action="store_true", help="skip schema validation")
    args = ap.parse_args(argv)

    fd = json.loads(Path(args.findings).read_text(encoding="utf-8"))
    tr = json.loads(Path(args.tests).read_text(encoding="utf-8"))

    if not args.no_validate:
        errors = validate(fd, "findingsFile", "findings.json") + validate(tr, "testResultsFile", "test_results.json")
        ids = [f["id"] for f in fd.get("findings", [])]
        errors += [f"findings.json: duplicate finding id {i}" for i, n in Counter(ids).items() if n > 1]
        known = set(ids)
        errors += [f"test_results.json: {t['test_id']} links unknown finding {t['linked_finding']}"
                   for t in tr.get("tests", []) if t.get("linked_finding") and t["linked_finding"] not in known]
        errors += [f"test_results.json: {t['test_id']} is NOT RUN without not_run_reason"
                   for t in tr.get("tests", []) if t.get("result") == "NOT RUN" and not t.get("not_run_reason")]
        if errors:
            print("Schema validation failed:", *errors, sep="\n  ", file=sys.stderr)
            return 2

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    sfx = suffix_for(fd["meta"])
    md_path = out / f"ANGRY_TESTER_REPORT{sfx}.md"
    xlsx_path = out / f"angry_tester_report{sfx}.xlsx"
    md_path.write_text(build_markdown(fd, tr), encoding="utf-8")
    build_excel(fd, tr, xlsx_path)

    sev = Counter(f["severity"] for f in fd["findings"])
    tc = test_counts(tr["tests"])
    print(f"Markdown: {md_path.resolve()}")
    print(f"Excel:    {xlsx_path.resolve()}")
    print("Findings: " + ", ".join(f"{s} {sev.get(s, 0)}" for s in SEVERITIES) + f" (total {len(fd['findings'])})")
    print(f"Tests:    PASS {tc['PASS']}, FAIL {tc['FAIL']}, NOT RUN {tc['NOT RUN']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
